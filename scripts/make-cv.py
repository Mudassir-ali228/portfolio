"""Build the CV .docx in the new template.

Follows the structure of the reference resume: Skills first, small-caps section
headings with a rule beneath, two-column entries (bold left / bold right, then
italic left / italic right), dot bullets under Experience and asterisk bullets
under Projects, Education last.

Client work is named generically and nothing is framed by industry — the same
rules the portfolio site follows.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "Mudassir_CV.docx"

SERIF = "Baskerville"
BODY_PT = 10.5
RIGHT_TAB = Inches(7.5)

# ------------------------------------------------------------------ content

NAME = "Mudassir Ali"
CONTACT = [
    "+92 309 0477778",
    "amudassir4321@gmail.com",
    "linkedin.com/in/mudassir-ali",
    "github.com/Mudassir-ali228",
]

SKILLS = [
    ("Languages", "JavaScript, TypeScript, Dart, Python, Java, C++, C#, SQL"),
    ("Web & Backend",
     "React.js, Next.js, Node.js, Express.js, NestJS, Spring Boot, REST API design, "
     "JWT, RBAC, Socket.IO, Prisma, Tailwind CSS"),
    ("Mobile & Games",
     "Flutter, React Native, Expo, Unity 6, native modules (Kotlin, Objective-C++), "
     "procedural generation, editor tooling"),
    ("Data & Infrastructure",
     "MySQL, PostgreSQL, MongoDB, ClickHouse, SQLite, Redis, Firebase, Docker, "
     "Nginx, systemd, Vercel, Netlify, Azure"),
    ("Tools & Testing",
     "Git/GitHub, Claude Code, Playwright, Selenium, Swagger/OpenAPI, Zod, PDFKit, JIRA"),
]

EXPERIENCE = [
    {
        "org": "SourceLabs", "dates": "6 Weeks",
        "role": "Web App Intern", "place": "Lahore, Pakistan",
        "bullets": [
            "Built and shipped full-stack features across **MongoDB, Express, React and Node.js** under senior-engineer review.",
            "Designed **REST** endpoints and component-based **React** interfaces with consistent client-side state management.",
        ],
    },
    {
        "org": "Coshi Infrastructure", "dates": "12 Weeks",
        "role": "Web App Intern", "place": "Causeway Bay, Hong Kong — Remote",
        "bullets": [
            "Delivered assigned modules across several web stacks, ramping quickly on unfamiliar codebases and conventions.",
            "Worked asynchronously with a distributed team across time zones, keeping every module on schedule.",
        ],
    },
]

PROJECTS = [
    {
        "name": "Company's Voice Portal",
        "stack": "React, Vite, Tailwind CSS, Node.js, Express, MySQL, JWT, PDFKit, REST API",
        "bullets": [
            "Built a **billing and invoicing platform** that syncs accounts, rate plans and usage from a provider API, prices every record against the client's own plan, and issues **PDF invoices** in the operator's existing layout.",
            "Implemented the **rating and tax engine** end to end — per-destination rounding, ordered tax chain, non-compounding arrears — plus usage import and SMTP invoice delivery.",
        ],
    },
    {
        "name": "RTP Logger",
        "stack": "NestJS, ClickHouse, Prisma, React, Material UI, Vector, Docker",
        "bullets": [
            "Built a self-hosted **log analytics platform** ingesting two high-volume feeds, one text and one binary, into **ClickHouse** through Vector and goflow2 — replacing an **Elasticsearch/Logstash** stack on hardware that could not afford it.",
            "Delivered a **NestJS** API with **JWT** auth and role-based access, and a React explorer with time-range search and dashboards; shipped both Docker Compose and native systemd deployments.",
        ],
    },
    {
        "name": "Company's CRM",
        "stack": "Express, Prisma, PostgreSQL, MySQL, Redis, Socket.IO, React, Redux Toolkit",
        "bullets": [
            "Contributed to a **multi-tenant SaaS CRM** integrating a third-party call-centre platform over its own database and event feed, with organisation-scoped access enforced in middleware.",
            "Built **live agent wallboards** on Redis-cached state pushed over **Socket.IO**, with Recharts supervisor dashboards reporting real-time queue and agent metrics.",
        ],
    },
    {
        "name": "I Drive — Open-World Driving Game",
        "stack": "Unity 6, C#, procedural generation",
        "bullets": [
            "Built a **1 km²** streamed open-world city with **no imported art**: roads, buildings, vehicles, characters, audio and the entire interface are generated from code at runtime, five materials for the whole city and **~2.5 ms** to rebuild a chunk.",
        ],
    },
    {
        "name": "I Run — Endless Runner",
        "stack": "Unity 6, C#, object pooling, editor tooling",
        "bullets": [
            "Built a three-lane endless runner whose generated layouts are **survivable by construction**, with **zero steady-state allocation** and an editor pass that bakes the runtime art into editable prefabs and scenes.",
        ],
    },
]

EDUCATION = [
    {
        "school": "COMSATS University Islamabad, Lahore Campus", "place": "Lahore, Pakistan",
        "degree": "Bachelor of Science in Software Engineering", "when": "Graduated Jun. 2026",
    },
    {
        "school": "Punjab Group of Colleges", "place": "Lahore, Pakistan",
        "degree": "Intermediate in Computer Science", "when": "Oct. 2020 – May 2022",
    },
]

# ------------------------------------------------------------------- helpers

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(0.45)
sec.bottom_margin = Inches(0.4)
sec.left_margin = sec.right_margin = Inches(0.5)

normal = doc.styles["Normal"]
normal.font.name = SERIF
normal.font.size = Pt(BODY_PT)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.0
rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rpr.set(qn(a), SERIF)


def para(space_before=0, space_after=0, left=0.0, hanging=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    if left:
        pf.left_indent = Inches(left)
    if hanging:
        pf.first_line_indent = Inches(-hanging)
    return p


def run(p, text, *, bold=False, italic=False, size=BODY_PT, caps=False, spacing=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.name = SERIF
    rf = r._element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), SERIF)
    if caps:
        sc = OxmlElement("w:smallCaps")
        r._element.get_or_add_rPr().append(sc)
    if spacing is not None:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))  # twentieths of a point
        r._element.get_or_add_rPr().append(sp)
    return r


def bottom_rule(p, size=6):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    bdr.append(bottom)
    pPr.append(bdr)


def section(title):
    p = para(space_before=5, space_after=2)
    run(p, title, bold=False, size=12, caps=True, spacing=0.6)
    bottom_rule(p)


def two_col(left_text, right_text, *, bold=False, italic=False, space_before=0):
    """Bold/italic pair with the right-hand item flushed to the margin."""
    p = para(space_before=space_before)
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    run(p, left_text, bold=bold, italic=italic)
    run(p, "\t")
    run(p, right_text, bold=bold, italic=italic)
    return p


MARK = __import__("re").compile(r"\*\*(.+?)\*\*")


def bullet(text, marker="•"):
    p = para(left=0.24, hanging=0.14, space_after=0.5)
    run(p, f"{marker} ")
    pos = 0
    for m in MARK.finditer(text):
        if m.start() > pos:
            run(p, text[pos:m.start()])
        run(p, m.group(1), bold=True)
        pos = m.end()
    if pos < len(text):
        run(p, text[pos:])
    return p


# -------------------------------------------------------------------- header

p = para()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, NAME, size=20, caps=True, spacing=1.2)

p = para(space_after=1)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, item in enumerate(CONTACT):
    if i:
        run(p, "  |  ", size=9)
    run(p, item, size=9)

# -------------------------------------------------------------------- skills

section("Skills")
for label, items in SKILLS:
    p = para(left=0.14, space_after=0.5)
    run(p, f"{label}: ", bold=True)
    run(p, items)

# ---------------------------------------------------------------- experience

section("Experience")
for i, job in enumerate(EXPERIENCE):
    two_col(job["org"], job["dates"], bold=True, space_before=3 if i else 1)
    two_col(job["role"], job["place"], italic=True)
    for b in job["bullets"]:
        bullet(b)

# ------------------------------------------------------------------ projects

section("Projects")
for i, proj in enumerate(PROJECTS):
    p = para(space_before=3 if i else 1)
    run(p, proj["name"], bold=True)
    run(p, " | ")
    run(p, proj["stack"], italic=True)
    for b in proj["bullets"]:
        bullet(b, marker="*")

# ----------------------------------------------------------------- education

section("Education")
for i, ed in enumerate(EDUCATION):
    two_col(ed["school"], ed["place"], bold=True, space_before=3 if i else 1)
    two_col(ed["degree"], ed["when"], italic=True)

doc.save(OUT)
print(f"wrote {OUT}")
